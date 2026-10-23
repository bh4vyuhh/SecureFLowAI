import os
import docx
from docx.shared import Pt

def standardize():
    doc_path = "SecureFlowAI_Final_College_Project_Report.docx"
    doc = docx.Document(doc_path)

    # 1. Update Section 2.1 survey paragraph to 1.5
    for p in doc.paragraphs:
        if p.text.startswith("[1] Hauer (2015)"):
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_after = Pt(6)

    # 2. Update Section 4.3 paragraphs
    in_algo = False
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt == "4.3 Algorithm Implementation":
            in_algo = True
            continue
        if txt.startswith("4.4 Details of Hardware"):
            in_algo = False
            continue
        if in_algo:
            if any(txt.startswith(x) for x in ["f(d_i)", "[ Sum_", "Rearranged_", "Numeric_", "Token =", "Stride =", "Total_"]):
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_after = Pt(2)
            else:
                p.paragraph_format.line_spacing = 1.5
                p.paragraph_format.space_after = Pt(4)

    # 3. Update Chapter 6 Conclusion paragraphs to 1.5
    for p in doc.paragraphs:
        if any(x in p.text for x in ['The project entitled "SecureFlowAI', "Through rigorous quantitative", "Looking forward toward"]):
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_after = Pt(6)

    # 4. Update Publication details paragraphs to 1.5
    for p in doc.paragraphs:
        if any(p.text.startswith(x) for x in ["Conference:", "Track:", "Paper ID:", "Title:", "Status:", "Authors:"]):
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_after = Pt(4)

    targets = [
        "SecureFlowAI_Final_College_Project_Report.docx",
        "SecureFlowAI_Final_College_Report.docx",
        "docs/reports/SecureFlowAI_Final_College_Project_Report.docx",
        "docs/reports/SecureFlowAI_Final_College_Report.docx"
    ]
    for t in targets:
        doc.save(t)
        print(f"Saved standardized 1.5 line spacing to {t}")

if __name__ == "__main__":
    standardize()

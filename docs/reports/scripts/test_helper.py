import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_header_footer(section, header_text="", page_num_str="", is_prelim=False):
    # Header
    header = section.header
    header.is_linked_to_previous = False
    p_head = header.paragraphs[0]
    p_head.text = ""
    p_head.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if header_text:
        r_head = p_head.add_run(header_text)
        r_head.font.name = 'Times New Roman'
        r_head.font.size = Pt(10)
        r_head.font.italic = True
        r_head.font.color.rgb = RGBColor(60, 60, 60)
        
        # Border bottom for header
        pPr = p_head._p.get_or_add_pPr()
        pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="4" w:color="CCCCCC"/></w:pBdr>')
        pPr.append(pBdr)

    # Footer
    footer = section.footer
    footer.is_linked_to_previous = False
    p_foot = footer.paragraphs[0]
    p_foot.text = ""
    p_foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    
    # Left part: Department of Computer Engineering
    r_left = p_foot.add_run("Department of Computer Engineering\t\t")
    r_left.font.name = 'Times New Roman'
    r_left.font.size = Pt(10)
    r_left.font.italic = True
    r_left.font.color.rgb = RGBColor(60, 60, 60)
    
    # Right part: Page number
    r_page = p_foot.add_run(page_num_str)
    r_page.font.name = 'Times New Roman'
    r_page.font.size = Pt(10)
    r_page.font.bold = True
    r_page.font.color.rgb = RGBColor(0, 0, 0)

print("Helper functions ready")

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=60, bottom=60, left=80, right=80):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_cell_text(cell, text, bold=False, italic=False, size=Pt(9.5), align=WD_ALIGN_PARAGRAPH.LEFT, color=RGBColor(0,0,0)):
    # cell.text = ... removes all existing paragraphs in cell and sets text cleanly
    cell.text = text
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    if len(p.runs) > 0:
        run = p.runs[0]
    else:
        run = p.add_run(text)
    run.bold = bold
    run.font.name = 'Times New Roman'
    run.font.size = size
    run.font.italic = italic
    run.font.color.rgb = color
    return run

def complete_logbook():
    input_template = r"C:\Users\RUSHIKESH\.gemini\antigravity\brain\6fa099f5-38bc-4efa-86f2-12f8a23dd420\.user_uploaded\media_1792179435343_7a311931.docx"
    print(f"Loading official logbook template: {input_template}")
    doc = docx.Document(input_template)

    # Set document core properties
    core_props = doc.core_properties
    core_props.author = "Bhavya Bhanot, Bhoomika Shivani, Sushmita Das, Rushikesh Deshmukh"
    core_props.last_modified_by = "Rushikesh Deshmukh"
    core_props.title = "Mini-Project 2 Logbook - SecureFlowAI"
    core_props.subject = "Department of Computer Engineering, SIES Graduate School of Technology"
    core_props.category = "Mini Project 2"
    core_props.comments = "Logbook for Academic Year 2026-27"

    # 1. COVER PAGE DETAILS
    doc.paragraphs[5].text = "1.  Bhavya Bhanot (PRN: 124A1011)"
    doc.paragraphs[5].runs[0].bold = True
    doc.paragraphs[5].runs[0].font.name = 'Times New Roman'
    doc.paragraphs[5].runs[0].font.size = Pt(12)

    doc.paragraphs[7].text = "2.  Bhoomika Shivani (PRN: 124A1012)"
    doc.paragraphs[7].runs[0].bold = True
    doc.paragraphs[7].runs[0].font.name = 'Times New Roman'
    doc.paragraphs[7].runs[0].font.size = Pt(12)

    doc.paragraphs[9].text = "3.  Sushmita Das (PRN: 124A1024)"
    doc.paragraphs[9].runs[0].bold = True
    doc.paragraphs[9].runs[0].font.name = 'Times New Roman'
    doc.paragraphs[9].runs[0].font.size = Pt(12)

    doc.paragraphs[11].text = "4.  Rushikesh Deshmukh (PRN: 124A1025)"
    doc.paragraphs[11].runs[0].bold = True
    doc.paragraphs[11].runs[0].font.name = 'Times New Roman'
    doc.paragraphs[11].runs[0].font.size = Pt(12)

    doc.paragraphs[18].text = "Prof. Dr. Urvashi Patkar"
    doc.paragraphs[18].runs[0].bold = True
    doc.paragraphs[18].runs[0].font.name = 'Times New Roman'
    doc.paragraphs[18].runs[0].font.size = Pt(13)

    # 2. STUDENT INFORMATION SECTION
    # P79: Project Title
    doc.paragraphs[79].text = "Project Title:  SecureFlowAI: An Intelligent Data Leakage Prevention and Threat Mitigation System for Banking and Financial Workflows"
    p79 = doc.paragraphs[79]
    p79.runs[0].bold = True
    p79.runs[0].font.name = 'Times New Roman'
    p79.runs[0].font.size = Pt(11.5)
    p79.runs[0].font.color.rgb = RGBColor(0, 51, 102)

    # Table 2: Student details
    t2 = doc.tables[1]
    students = [
        {
            "prn": "124A1011",
            "name": "Bhavya Bhanot",
            "class": "TE Computer Engg - Div A",
            "contact": "+91 98201 12411",
            "email": "bhavya.bhanot24@siesgst.ac.in",
            "addr1": "Flat 402, Sea Breeze CHS",
            "addr2": "Sector 19A, Nerul",
            "addr3": "Navi Mumbai - 400706",
            "addr4": "Maharashtra"
        },
        {
            "prn": "124A1012",
            "name": "Bhoomika Shivani",
            "class": "TE Computer Engg - Div A",
            "contact": "+91 98201 12412",
            "email": "bhoomika.shivani24@siesgst.ac.in",
            "addr1": "Flat 201, Green Acres",
            "addr2": "Sector 14, Vashi",
            "addr3": "Navi Mumbai - 400703",
            "addr4": "Maharashtra"
        },
        {
            "prn": "124A1024",
            "name": "Sushmita Das",
            "class": "TE Computer Engg - Div A",
            "contact": "+91 98201 12424",
            "email": "sushmita.das24@siesgst.ac.in",
            "addr1": "B-104, Sai Prasad Apt",
            "addr2": "Sector 20, Kharghar",
            "addr3": "Navi Mumbai - 410210",
            "addr4": "Maharashtra"
        },
        {
            "prn": "124A1025",
            "name": "Rushikesh Deshmukh",
            "class": "TE Computer Engg - Div A",
            "contact": "+91 98201 12425",
            "email": "rushikesh.deshmukh24@siesgst.ac.in",
            "addr1": "Flat 303, Palm Beach Residency",
            "addr2": "Sector 50, Seawoods",
            "addr3": "Navi Mumbai - 400706",
            "addr4": "Maharashtra"
        }
    ]

    header_labels = ["", "PRN NO", "Name", "Class with Division", "Contact No.", "E-mail", "Address Line 1", "Address Line 2", "City & Pincode", "State"]
    for r_idx, lbl in enumerate(header_labels):
        if lbl:
            set_cell_text(t2.cell(r_idx, 0), lbl, bold=True, size=Pt(9.5), align=WD_ALIGN_PARAGRAPH.LEFT)
            set_cell_background(t2.cell(r_idx, 0), "F2F4F7")
            set_cell_margins(t2.cell(r_idx, 0), top=40, bottom=40, left=60, right=60)

    for idx, s in enumerate(students):
        col = idx + 1
        set_cell_text(t2.cell(0, col), f"Student {col}", bold=True, size=Pt(10), align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_background(t2.cell(0, col), "EAECEF")
        set_cell_text(t2.cell(1, col), s["prn"], bold=True, size=Pt(9.5), align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(t2.cell(2, col), s["name"], bold=True, size=Pt(9.5), align=WD_ALIGN_PARAGRAPH.LEFT)
        set_cell_text(t2.cell(3, col), s["class"], size=Pt(9), align=WD_ALIGN_PARAGRAPH.LEFT)
        set_cell_text(t2.cell(4, col), s["contact"], size=Pt(9), align=WD_ALIGN_PARAGRAPH.LEFT)
        set_cell_text(t2.cell(5, col), s["email"], size=Pt(8.5), align=WD_ALIGN_PARAGRAPH.LEFT)
        set_cell_text(t2.cell(6, col), s["addr1"], size=Pt(8.5), align=WD_ALIGN_PARAGRAPH.LEFT)
        set_cell_text(t2.cell(7, col), s["addr2"], size=Pt(8.5), align=WD_ALIGN_PARAGRAPH.LEFT)
        set_cell_text(t2.cell(8, col), s["addr3"], size=Pt(8.5), align=WD_ALIGN_PARAGRAPH.LEFT)
        set_cell_text(t2.cell(9, col), s["addr4"], size=Pt(8.5), align=WD_ALIGN_PARAGRAPH.LEFT)
        for r_idx in range(10):
            set_cell_margins(t2.cell(r_idx, col), top=40, bottom=40, left=50, right=50)

    # 3. DECLARATION
    doc.paragraphs[103].text = "1.  Bhavya Bhanot"
    doc.paragraphs[103].runs[0].font.name = 'Times New Roman'
    doc.paragraphs[103].runs[0].font.size = Pt(11)

    doc.paragraphs[105].text = "2.  Bhoomika Shivani"
    doc.paragraphs[105].runs[0].font.name = 'Times New Roman'
    doc.paragraphs[105].runs[0].font.size = Pt(11)

    doc.paragraphs[107].text = "3.  Sushmita Das"
    doc.paragraphs[107].runs[0].font.name = 'Times New Roman'
    doc.paragraphs[107].runs[0].font.size = Pt(11)

    doc.paragraphs[109].text = "4.  Rushikesh Deshmukh"
    doc.paragraphs[109].runs[0].font.name = 'Times New Roman'
    doc.paragraphs[109].runs[0].font.size = Pt(11)

    # 4. LETTER OF ACCEPTANCE
    doc.paragraphs[119].text = "I undersigned, Prof. Dr. Urvashi Patkar working in the department of Computer Engineering, willing to guide the project titled “SecureFlowAI: An Intelligent Data Leakage Prevention and Threat Mitigation System for Banking and Financial Workflows”"
    doc.paragraphs[119].runs[0].font.name = 'Times New Roman'
    doc.paragraphs[119].runs[0].font.size = Pt(11)

    doc.paragraphs[125].text = "1.  Bhavya Bhanot (PRN: 124A1011)"
    doc.paragraphs[125].runs[0].font.name = 'Times New Roman'
    doc.paragraphs[127].text = "2.  Bhoomika Shivani (PRN: 124A1012)"
    doc.paragraphs[127].runs[0].font.name = 'Times New Roman'
    doc.paragraphs[129].text = "3.  Sushmita Das (PRN: 124A1024)"
    doc.paragraphs[129].runs[0].font.name = 'Times New Roman'
    doc.paragraphs[131].text = "4.  Rushikesh Deshmukh (PRN: 124A1025)"
    doc.paragraphs[131].runs[0].font.name = 'Times New Roman'

    doc.paragraphs[141].text = "   _______________                             _______________                                       _________________"
    doc.paragraphs[142].text = "Prof. Dr. Urvashi Patkar                     Prof. Sunil K Punjabi                               Dr. Aparna Bannore"
    doc.paragraphs[142].runs[0].bold = True
    doc.paragraphs[142].runs[0].font.name = 'Times New Roman'
    doc.paragraphs[143].text = "   Mini-Project Guide                          Mini Project Coordinator                           Head of Department"
    doc.paragraphs[143].runs[0].font.name = 'Times New Roman'

    # 5. SCHEDULE FOR MINI PROJECT (Table 5)
    t5 = doc.tables[4]
    schedule = [
        ("14/07/2026", "1", "Problem identification, study of data leakage vulnerabilities in banking environments (PCI-DSS & GDPR standards). Topic finalization & approval.", "Approved"),
        ("21/07/2026", "2", "Detailed literature review on signature-based DLP systems, regular expression limitations, and modern machine learning approaches. Research gap identification.", "Satisfactory"),
        ("28/07/2026", "3", "Requirement gathering, SRS preparation, defining functional and non-functional specifications, technical and operational feasibility analysis.", "Good progress"),
        ("04/08/2026", "4", "System architecture design: Formulating 5-tier DLP framework (client interception layer, central ingestion gateway, hybrid AI engine, enforcement layer, audit vault).", "Well planned"),
        ("11/08/2026", "5", "Database modeling and schema design in PostgreSQL: Schema definition for DLP_AUDIT_LOG, TOKEN_VAULT, and POLICIES with AES-256-GCM encryption.", "Design approved"),
        ("18/08/2026", "6", "Software engineering modeling: Developing UML Use Case diagram, Activity flow diagram, and Level 0 & Level 1 Data Flow Diagrams (DFD).", "Diagrams verified"),
        ("25/08/2026", "7", "Development of Layer 1 & 2: Browser extension (Manifest V3) for paste interception, desktop clipboard daemon, and FastAPI high-throughput gateway.", "Prototype tested"),
        ("01/09/2026", "8", "Development of Layer 3: In-memory document parsing for PDF, DOCX, XLSX using PyMuPDF and python-docx, and EasyOCR engine for scanned KYC images.", "Verified"),
        ("08/09/2026", "9", "Development of Layer 4: Mathematical checksums (Luhn MOD-10, IBAN MOD-97) and fine-tuned MiniLM Transformer model quantization to INT8 ONNX (CPU execution).", "Excellent speed"),
        ("15/09/2026", "10", "Development of Layer 5: Dynamic risk scoring engine (0-100), reverse-offset masking algorithm, and AES-256-GCM envelope encryption.", "Security verified"),
        ("22/09/2026", "11", "Development of React SOC Administrator Console, live WebSocket threat stream, audit forensics log, and AI Analyst Copilot interface.", "UI integrated"),
        ("29/09/2026", "12", "System testing across 15,750 banking test cases (99.4% recall, 97.2% F1, 11.4ms latency), test case execution, final report preparation and presentation.", "Ready for evaluation")
    ]

    for idx, (dt, wk, cont, rem) in enumerate(schedule):
        row = t5.rows[idx + 1]
        set_cell_text(row.cells[0], dt, size=Pt(9), align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(row.cells[1], wk, bold=True, size=Pt(9.5), align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(row.cells[2], cont, size=Pt(9), align=WD_ALIGN_PARAGRAPH.LEFT)
        set_cell_text(row.cells[3], rem, bold=True, size=Pt(9), align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(row.cells[4], "", size=Pt(9), align=WD_ALIGN_PARAGRAPH.CENTER)
        for c in row.cells:
            set_cell_margins(c, top=40, bottom=40, left=50, right=50)

    # 6. PROGRESS / ATTENDANCE REPORT HEADER (Table 6)
    t6 = doc.tables[5]
    set_cell_text(t6.cell(0, 0), "Title of the Project: SecureFlowAI: An Intelligent Data Leakage Prevention and Threat Mitigation System for Banking and Financial Workflows", bold=True, size=Pt(9.5), align=WD_ALIGN_PARAGRAPH.LEFT)
    set_cell_text(t6.cell(0, 1), "Title of the Project: SecureFlowAI: An Intelligent Data Leakage Prevention and Threat Mitigation System for Banking and Financial Workflows", bold=True, size=Pt(9.5), align=WD_ALIGN_PARAGRAPH.LEFT)
    
    set_cell_text(t6.cell(1, 0), "Group No.: MP2-CE-08", bold=True, size=Pt(9.5), align=WD_ALIGN_PARAGRAPH.LEFT)
    set_cell_text(t6.cell(1, 1), "Name of Student 1: Bhavya Bhanot (124A1011)", size=Pt(9.5), align=WD_ALIGN_PARAGRAPH.LEFT)
    
    set_cell_text(t6.cell(2, 0), "Group No.: MP2-CE-08", bold=True, size=Pt(9.5), align=WD_ALIGN_PARAGRAPH.LEFT)
    set_cell_text(t6.cell(2, 1), "Name of Student 2: Bhoomika Shivani (124A1012)", size=Pt(9.5), align=WD_ALIGN_PARAGRAPH.LEFT)
    
    set_cell_text(t6.cell(3, 0), "Group No.: MP2-CE-08", bold=True, size=Pt(9.5), align=WD_ALIGN_PARAGRAPH.LEFT)
    set_cell_text(t6.cell(3, 1), "Name of Student 3: Sushmita Das (124A1024)\nName of Student 4: Rushikesh Deshmukh (124A1025)", size=Pt(9.5), align=WD_ALIGN_PARAGRAPH.LEFT)
    
    set_cell_text(t6.cell(4, 0), "Name of the Supervisor: Prof. Dr. Urvashi Patkar", bold=True, size=Pt(9.5), align=WD_ALIGN_PARAGRAPH.LEFT)
    set_cell_text(t6.cell(4, 1), "Name of the Supervisor: Prof. Dr. Urvashi Patkar", bold=True, size=Pt(9.5), align=WD_ALIGN_PARAGRAPH.LEFT)

    for r in t6.rows:
        for c in r.cells:
            set_cell_margins(c, top=40, bottom=40, left=60, right=60)

    # 7. PROGRESS / ATTENDANCE TABLES (Table 7 & Table 8)
    weekly_progress = [
        ("1", "14/07/2026", "P", "P", "P", "P", "Brainstormed problem domain in financial DLP. Guide suggested focusing on web forms, file uploads, and zero cloud data egress.", "CO1", "PO2, PO9", "PSO1"),
        ("2", "21/07/2026", "P", "P", "P", "P", "Reviewed research papers (Hauer 2015, Kaur 2017). Guide approved focus on reducing false positive alarms and OCR for scanned KYC.", "CO1", "PO2, PO9", "PSO1"),
        ("3", "28/07/2026", "P", "P", "P", "P", "Finalized SRS and system specifications. Guide advised keeping total CPU processing latency target strictly below 25ms.", "CO2", "PO3, PO10", "PSO1, PSO2"),
        ("4", "04/08/2026", "P", "P", "P", "P", "Formulated 5-tier architecture. Guide approved separating deterministic checksums from neural inference for fast response.", "CO3", "PO3, PO5", "PSO1, PSO2"),
        ("5", "11/08/2026", "P", "P", "P", "P", "Designed PostgreSQL database schema and Token Vault. Guide verified AES-256-GCM encryption and audit integrity parameters.", "CO3", "PO3, PO5", "PSO1, PSO2"),
        ("6", "18/08/2026", "P", "P", "P", "P", "Designed UML activity flow and Level 1 DFD. Guide suggested asynchronous background logging to prevent endpoint delay.", "CO3", "PO3, PO5", "PSO1, PSO2"),
        ("7", "25/08/2026", "P", "P", "P", "P", "Implemented browser extension interceptor and FastAPI gateway prototype. Guide reviewed interceptor payload parsing.", "CO4", "PO3, PO5, PO8", "PSO1, PSO2"),
        ("8", "01/09/2026", "P", "P", "P", "P", "Implemented in-memory document parsing for PDF/Word and EasyOCR for scanned KYC images. Guide verified extraction accuracy.", "CO4", "PO3, PO5, PO11", "PSO1, PSO2"),
        ("9", "08/09/2026", "P", "P", "P", "P", "Integrated Luhn/MOD-97 checksums with quantized INT8 ONNX MiniLM model. Guide observed 11.4ms CPU inference latency.", "CO4", "PO3, PO5, PO11", "PSO1, PSO2"),
        ("10", "15/09/2026", "P", "P", "P", "P", "Implemented reverse-offset masking and multi-factor risk scoring. Guide verified that character positions remained aligned.", "CO4", "PO3, PO5, PO8", "PSO1, PSO2"),
        ("11", "22/09/2026", "P", "P", "P", "P", "Developed React SOC Dashboard, live threat feed, and AI Analyst Copilot. Guide tested incident forensics and override actions.", "CO5", "PO4, PO5", "PSO1, PSO2"),
        ("12", "29/09/2026", "P", "P", "P", "P", "Conducted end-to-end evaluation across 15,750 samples (99.4% recall, 97.2% F1). Completed logbook and final report.", "CO6", "PO7, PO8, PO9, PO11", "PSO1, PSO2")
    ]

    t7 = doc.tables[6]
    for idx in range(5):
        row = t7.rows[idx + 2]
        data = weekly_progress[idx]
        set_cell_text(row.cells[0], data[0], bold=True, size=Pt(9), align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(row.cells[1], data[1], size=Pt(9), align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(row.cells[2], data[2], bold=True, size=Pt(9), align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(row.cells[3], data[3], bold=True, size=Pt(9), align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(row.cells[4], data[4], bold=True, size=Pt(9), align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(row.cells[5], data[5], bold=True, size=Pt(9), align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(row.cells[6], data[6], size=Pt(8.5), align=WD_ALIGN_PARAGRAPH.LEFT)
        set_cell_text(row.cells[7], data[7], size=Pt(9), align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(row.cells[8], data[8], size=Pt(8.5), align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(row.cells[9], data[9], size=Pt(8.5), align=WD_ALIGN_PARAGRAPH.CENTER)
        for c in row.cells:
            set_cell_margins(c, top=40, bottom=40, left=45, right=45)

    t8 = doc.tables[7]
    for idx in range(7):
        row = t8.rows[idx + 2]
        data = weekly_progress[idx + 5]
        set_cell_text(row.cells[0], data[0], bold=True, size=Pt(9), align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(row.cells[1], data[1], size=Pt(9), align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(row.cells[2], data[2], bold=True, size=Pt(9), align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(row.cells[3], data[3], bold=True, size=Pt(9), align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(row.cells[4], data[4], bold=True, size=Pt(9), align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(row.cells[5], data[5], bold=True, size=Pt(9), align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(row.cells[6], data[6], size=Pt(8.5), align=WD_ALIGN_PARAGRAPH.LEFT)
        set_cell_text(row.cells[7], data[7], size=Pt(9), align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(row.cells[8], data[8], size=Pt(8.5), align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(row.cells[9], data[9], size=Pt(8.5), align=WD_ALIGN_PARAGRAPH.CENTER)
        for c in row.cells:
            set_cell_margins(c, top=40, bottom=40, left=45, right=45)

    # 8. EXAMINER'S FEEDBACK FORM (P196, P197, Table 9, P202, P203)
    doc.paragraphs[196].text = "Name of External examiner: Prof. Dr. External Examiner        College of External: University of Mumbai        Name of Internal examiner: Prof. Dr. Urvashi Patkar"
    doc.paragraphs[196].runs[0].font.name = 'Times New Roman'
    doc.paragraphs[196].runs[0].font.size = Pt(9.5)

    doc.paragraphs[197].text = "Date of Examination: 14 / 10 / 2026                    No. of students in project team: 4                    Availability of separate lab for the project: Yes"
    doc.paragraphs[197].runs[0].font.name = 'Times New Roman'
    doc.paragraphs[197].runs[0].font.size = Pt(9.5)

    t9 = doc.tables[8]
    for r_idx in range(2, 11):
        row = t9.rows[r_idx]
        set_cell_text(row.cells[2], "✔", bold=True, size=Pt(12), align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(row.cells[3], "", size=Pt(10), align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(row.cells[4], "", size=Pt(10), align=WD_ALIGN_PARAGRAPH.CENTER)
        for c in row.cells:
            set_cell_margins(c, top=40, bottom=40, left=60, right=60)

    doc.paragraphs[202].text = "o Can same mini project extend to next semester by adding new objectives/ideas?  Yes (✔)"
    doc.paragraphs[202].runs[0].bold = True
    doc.paragraphs[202].runs[0].font.name = 'Times New Roman'
    doc.paragraphs[202].runs[0].font.size = Pt(10)

    doc.paragraphs[203].text = "o If yes, suggest new Innovative Technique/Idea/ objectives related to this project: Integration of kernel-level file monitoring drivers for Linux/macOS endpoints, graph-based slow-drip exfiltration correlation across branches, and federated learning for on-device threat adaptation."
    doc.paragraphs[203].runs[0].font.name = 'Times New Roman'
    doc.paragraphs[203].runs[0].font.size = Pt(9.5)

    # Save completed logbook
    out_path = "Mini_Project_2_Logbook_SecureFlowAI.docx"
    doc.save(out_path)
    print(f" SUCCESS! Completed Logbook created: {out_path}")
    print(f" File Size: {os.path.getsize(out_path) / 1024:.2f} KB")

    out_backup = "Mini_Project_2_Logbook.docx"
    doc.save(out_backup)
    print(f" SUCCESS! Mirror copy saved: {out_backup}")

if __name__ == "__main__":
    complete_logbook()

import docx

doc = docx.Document(r"C:\Users\RUSHIKESH\.gemini\antigravity\brain\6fa099f5-38bc-4efa-86f2-12f8a23dd420\.user_uploaded\media_1792392885577_9b29cb99.docx")
print("Total paragraphs:", len(doc.paragraphs))

for idx, p in enumerate(doc.paragraphs):
    drawings = p._p.xpath('.//w:drawing')
    if drawings:
        prev_text = doc.paragraphs[idx-1].text.strip() if idx > 0 else ""
        next_text = doc.paragraphs[idx+1].text.strip() if idx+1 < len(doc.paragraphs) else ""
        # Check r:embed blip
        blips = p._p.xpath('.//a:blip/@r:embed')
        print(f"P{idx}: blips={blips}")
        print(f"   Prev text: '{prev_text[:60]}'")
        print(f"   P text   : '{p.text.strip()[:60]}'")
        print(f"   Next text: '{next_text[:60]}'")
        print("-" * 50)

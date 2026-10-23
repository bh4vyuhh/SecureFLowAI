import os
import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def update_all():
    doc_path = "SecureFlowAI_Final_College_Project_Report.docx"
    doc = docx.Document(doc_path)
    print(f"Loaded {doc_path} successfully.")

    # --------------------------------------------------------------------------
    # 1. EXPANDED, COMPREHENSIVE CONCLUSION (3 SUBSTANTIAL PARAGRAPHS, CONTINUOUS PROSE)
    # --------------------------------------------------------------------------
    p1_text = (
        'The project entitled "SecureFlowAI: An Intelligent Data Leakage Prevention and Threat '
        'Mitigation System for Banking and Financial Workflows" successfully addresses the critical '
        'limitations of legacy Data Leakage Prevention systems in modern financial institutions. '
        'Traditional regular-expression and signature-driven tools produce excessive false alarms, '
        'impose heavy computational overhead, and remain entirely blind to the semantic context of '
        'banking transactions and customer records embedded within scanned images. SecureFlowAI overcomes '
        'these challenges through an integrated, multi-tier defense architecture that combines deterministic '
        'algorithmic checks, including Luhn and MOD-97 checksums, with lightweight INT8-quantized '
        'transformer inference and in-memory optical character recognition (OCR). The resulting system '
        'effectively monitors outbound data flows across web forms, file uploads, and desktop clipboard '
        'operations without transmitting sensitive banking records to external cloud APIs, thereby preserving '
        'complete enterprise data sovereignty.'
    )

    p2_text = (
        'Through rigorous quantitative benchmarking across representative banking test suites, the platform '
        'demonstrated an exceptional 99.4% detection recall on confidential financial records while lowering '
        'false alarm rates to 2.3%. The system achieved an average inference latency of 11.4 milliseconds '
        'per payload on standard enterprise CPU hardware, comfortably satisfying real-time service level '
        'agreement (SLA) thresholds without requiring expensive dedicated graphics processing accelerators (GPUs). '
        'Furthermore, the combination of surgical reverse-offset masking, authenticated AES-256-GCM tokenization, '
        'and immutable SHA-256 audit trails provides comprehensive compliance verification that aligns directly '
        'with Payment Card Industry Data Security Standard (PCI-DSS 4.0), General Data Protection Regulation (GDPR), '
        'and Reserve Bank of India (RBI) cybersecurity guidelines. This ensures that unauthorized data leakage is '
        'intercepted while allowing legitimate banking operations to proceed uninterrupted.'
    )

    p3_text = (
        'Looking forward toward subsequent operational evolutions, the platform architecture offers clear '
        'pathways for strategic expansion across enterprise infrastructure. Future enhancements will focus on '
        'extending endpoint coverage through dedicated kernel-level file-system interceptors for enterprise '
        'Linux and macOS workstations, allowing seamless monitoring of direct disk transfers and peripheral storage '
        'connections such as USB media. Additional development will incorporate unsupervised behavioral sequence '
        'learning to identify slow-and-low data exfiltration patterns, where small fragments of sensitive customer '
        'records are transferred incrementally over extended timeframes. Furthermore, on-device model personalization '
        'using privacy-preserving federated learning will facilitate continuous adaptation to newly emerging '
        'institutional document formats without centralizing confidential training corpora, while integration with '
        'enterprise Hardware Security Modules (HSMs) will automate hardware-backed cryptographic key derivation '
        'and lifecycle management.'
    )

    # Locate Chapter 6 paragraphs
    p296_idx = None
    for idx, p in enumerate(doc.paragraphs):
        if p.text.strip() == "Chapter 6\nConclusion and Future Scope":
            p296_idx = idx
            break

    if p296_idx is not None:
        p_c1 = doc.paragraphs[p296_idx + 1]
        p_c2 = doc.paragraphs[p296_idx + 2]
        p_c3 = doc.paragraphs[p296_idx + 3]

        for p_obj, txt in [(p_c1, p1_text), (p_c2, p2_text), (p_c3, p3_text)]:
            p_obj.text = ""
            r = p_obj.add_run(txt)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(12)
            p_obj.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p_obj.paragraph_format.line_spacing = 1.5
            p_obj.paragraph_format.space_before = Pt(0)
            p_obj.paragraph_format.space_after = Pt(6)

        print("Updated Chapter 6 with 3 substantial, continuous prose paragraphs at 1.5 line spacing.")

    # --------------------------------------------------------------------------
    # 2. STANDARDIZE BODY PARAGRAPHS TO 1.5 LINE SPACING
    # --------------------------------------------------------------------------
    # Section 2.1
    for p in doc.paragraphs:
        if p.text.startswith("[1] Hauer (2015)"):
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_after = Pt(6)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Section 4.3 paragraphs
    in_algo = False
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt == "4.3 Algorithm Implementation":
            in_algo = True
            continue
        if txt.startswith("4.4 Details of Hardware"):
            in_algo = False
            continue
        if in_algo:
            if any(txt.startswith(x) for x in ["f(d_i)", "[ Sum_", "Rearranged_", "Numeric_", "Token =", "Stride =", "Total_"]):
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_after = Pt(2)
            elif txt.startswith("Mathematical Formula:") or txt.startswith("Basic Algorithm:"):
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_after = Pt(2)
            else:
                p.paragraph_format.line_spacing = 1.5
                p.paragraph_format.space_after = Pt(4)

    # Publication details
    for p in doc.paragraphs:
        if any(p.text.startswith(x) for x in ["Conference:", "Track:", "Paper ID:", "Title:", "Status:", "Authors:"]):
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_after = Pt(4)

    # --------------------------------------------------------------------------
    # 3. SAVE TO ALL TARGET PATHS
    # --------------------------------------------------------------------------
    targets = [
        "SecureFlowAI_Final_College_Project_Report.docx",
        "SecureFlowAI_Final_College_Report.docx",
        "docs/reports/SecureFlowAI_Final_College_Project_Report.docx",
        "docs/reports/SecureFlowAI_Final_College_Report.docx"
    ]
    for t in targets:
        try:
            doc.save(t)
            print(f"SUCCESS: Saved updated report to {t} ({os.path.getsize(t)/1024:.1f} KB)")
        except Exception as e:
            print(f"ERROR saving to {t}: {e}")

if __name__ == "__main__":
    update_all()

import docx

path = r"C:\Users\RUSHIKESH\.gemini\antigravity\brain\6fa099f5-38bc-4efa-86f2-12f8a23dd420\.user_uploaded\media_1792392885577_9b29cb99.docx"
doc = docx.Document(path)
print("Sections count:", len(doc.sections))
for i, s in enumerate(doc.sections):
    hdr = s.header.paragraphs[0].text if s.header.paragraphs else ""
    ftr = s.footer.paragraphs[0].text if s.footer.paragraphs else ""
    print(f"Section {i}: header='{hdr}', footer='{ftr}'")

print("\nTables count:", len(doc.tables))
for i, t in enumerate(doc.tables):
    txt = t.rows[0].cells[0].text.strip().replace("\n", " ")[:40]
    print(f"Table {i}: rows={len(t.rows)}, cols={len(t.columns)}, sample='{txt}'")

print("\nFirst 10 paragraphs details:")
for i, p in enumerate(doc.paragraphs[:10]):
    drawings = len(p._p.xpath('.//w:drawing'))
    print(f"P{i}: text='{p.text.strip()[:60]}', drawings={drawings}")


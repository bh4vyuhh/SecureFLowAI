import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def update_report():
    doc_path = "SecureFlowAI_Final_College_Project_Report.docx"
    doc = docx.Document(doc_path)
    print(f"Loaded {doc_path} with {len(doc.paragraphs)} paragraphs and {len(doc.tables)} tables.")

    # =========================================================================
    # 1. UPDATE CHAPTER 2: LITERATURE SURVEY WITH INDICES
    # =========================================================================
    # Find paragraph "2.1 Survey of Existing System"
    p_survey_idx = None
    for idx, p in enumerate(doc.paragraphs):
        if p.text.strip() == "2.1 Survey of Existing System":
            p_survey_idx = idx
            break

    if p_survey_idx is not None:
        p_survey = doc.paragraphs[p_survey_idx + 1]
        p_survey.text = ""
        
        survey_runs = [
            ("[1] Hauer (2015)", True),
            (", published in IEEE Access and indexed in IEEE Xplore, Scopus, and Web of Science, examined real-world enterprise data breaches and found that simple signature-based tools miss many common leaks because their rules are too rigid. The study noted that signature systems struggle with dynamic data variations and fail to provide an automated machine learning classification pipeline.\n\n", False),
            ("[2] Kaur et al. (2017)", True),
            (", published in CS & IT-CSCP and indexed in Scopus and Google Scholar, conducted a comparative survey of signature, statistical, and specification-based DLP techniques. The authors noted that manual regular expression rules require constant tuning, suffer from high maintenance overhead, and completely miss text hidden inside image files or screenshots due to lack of optical character recognition.\n\n", False),
            ("[3] Ghouse et al. (2019)", True),
            (", published in IEEE ICAC3 and indexed in IEEE Xplore and Scopus, proposed a machine learning approach to classify and selectively encrypt sensitive data in transit within email communication. While effective for simple email text bodies, the system did not support multi-format file attachments, scanned image OCR, dynamic risk scoring, or a centralized security operations audit dashboard.\n\n", False),
            ("[4] Abid (2024)", True),
            (", published in Global Journal of Computer Science and Technology and indexed in CrossRef and Google Scholar, reviewed modern AI, machine learning, and cloud-native DLP frameworks within Zero Trust security architectures. The author highlighted that many proposed cloud-based systems remain theoretical, incur significant processing latency, and pose data privacy risks when transmitting corporate data to external cloud services.", False)
        ]
        
        for txt, is_bold in survey_runs:
            r = p_survey.add_run(txt)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(12)
            r.bold = is_bold
        
        p_survey.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_survey.paragraph_format.line_spacing = 1.15
        p_survey.paragraph_format.space_after = Pt(8)
        print("Updated Section 2.1 with verified publication indices.")

    # Update Table 2.1 (doc.tables[3])
    if len(doc.tables) > 3:
        t = doc.tables[3]
        table_data = [
            ["Index", "Author(s) & Year", "Method Used", "Key Findings & Identified Gaps"],
            ["[1]", "Hauer (2015)\nIEEE Access\n[IEEE Xplore, Scopus, WoS]", "Content-aware signature DLP analysis on enterprise breach records.", "Signature systems fail on complex leak scenarios; taxonomy is too rigid; lacks an implemented ML pipeline."],
            ["[2]", "Kaur et al. (2017)\nCS & IT-CSCP\n[Scopus, Google Scholar]", "Comparative survey of signature, statistical, and specification-based DLP.", "Manual regex rules require constant tuning; image files and screenshots are not scanned; cross-OS limitations."],
            ["[3]", "Ghouse et al. (2019)\nIEEE ICAC3\n[IEEE Xplore, Scopus]", "Machine learning text classification with selective data-in-transit encryption.", "Limited to email text bodies only; no file parsing, no OCR, no risk scoring engine, no audit dashboard."],
            ["[4]", "Abid (2024)\nGlobal Journal\n[CrossRef, Google Scholar]", "Review of AI/ML, cloud-native DLP, and Zero Trust security architectures.", "Review-level paper; lacks an implemented working pipeline, concrete risk formulas, and empirical verification."]
        ]
        
        for r_idx, row_values in enumerate(table_data):
            row = t.rows[r_idx]
            for c_idx, val in enumerate(row_values):
                cell = row.cells[c_idx]
                cell.text = ""
                p = cell.paragraphs[0]
                r = p.add_run(val)
                r.font.name = 'Times New Roman'
                r.font.size = Pt(10)
                if r_idx == 0:
                    r.bold = True
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    if c_idx == 0:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        r.bold = True
                    else:
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_after = Pt(2)
        print("Updated Table 2.1 with Index and indexing sources.")

    # =========================================================================
    # 2. UPDATE SECTION 4.3: ALGORITHM IMPLEMENTATION WITH CONTENT, FORMULAS & BASIC ALGOS
    # =========================================================================
    # Find paragraph "4.3 Algorithm Implementation"
    p_algo_idx = None
    for idx, p in enumerate(doc.paragraphs):
        if p.text.strip() == "4.3 Algorithm Implementation":
            p_algo_idx = idx
            break

    # Find the next section header "4.4 Details of Hardware & Software"
    p_hardware_idx = None
    for idx, p in enumerate(doc.paragraphs):
        if p.text.strip().startswith("4.4 Details of Hardware"):
            p_hardware_idx = idx
            break

    print(f"Algorithm section spans from P{p_algo_idx} to P{p_hardware_idx}")

    if p_algo_idx is not None and p_hardware_idx is not None:
        # Clear paragraphs between p_algo_idx and p_hardware_idx (exclusive)
        # Note: we modify in-place or remove paragraphs from p_algo_idx + 1 up to p_hardware_idx - 1
        # In python-docx, deleting paragraphs is done via XML p.getparent().remove(p._p)
        paragraphs_to_remove = doc.paragraphs[p_algo_idx + 1 : p_hardware_idx]
        for p in paragraphs_to_remove:
            p._p.getparent().remove(p._p)
        print(f"Cleared {len(paragraphs_to_remove)} old paragraphs in Section 4.3.")

        # Re-fetch reference paragraph for insertion
        ref_p = doc.paragraphs[p_algo_idx]

        # Content to insert
        sections = [
            {
                "title": "4.3.1 Luhn Algorithm for Credit Card Validation",
                "desc": "The Luhn algorithm (ISO/IEC 7812-1 MOD-10) checks whether a 16-digit primary account number (PAN) is mathematically valid. It doubles every second digit from right to left, subtracts 9 from any product exceeding 9, and computes the cumulative sum. If the total sum modulo 10 equals zero, the card number is mathematically authentic. This instantly filters out non-card 16-digit tracking codes in less than a microsecond without requiring neural network processing.",
                "formula_header": "Mathematical Formula:",
                "formula_lines": [
                    "f(d_i) = d_i  (if position i is odd)",
                    "f(d_i) = 2 * d_i  (if position i is even and 2 * d_i <= 9)",
                    "f(d_i) = (2 * d_i) - 9  (if position i is even and 2 * d_i > 9)",
                    "[ Sum_{i=1}^{n} f(d_i) ] mod 10 = 0"
                ],
                "algo_header": "Basic Algorithm:",
                "steps": [
                    "1. Extract candidate 16-digit numerical sequences from the incoming text payload.",
                    "2. Remove all spaces, hyphens, and non-digit formatting characters.",
                    "3. Transform each digit according to the odd/even position doubling rule.",
                    "4. Calculate the cumulative sum of all transformed digits and evaluate modulo 10.",
                    "5. If the remainder equals 0, confirm the string as a valid credit card; otherwise, mark it as non-sensitive."
                ]
            },
            {
                "title": "4.3.2 ISO 7064 MOD-97 Algorithm for IBAN Validation",
                "desc": "The MOD-97 algorithm verifies International Bank Account Numbers (IBAN) across global financial networks. The first four characters (country code and check digits) are moved to the end of the string. Letters A through Z are converted to numerical values from 10 to 35. The resulting integer is divided by 97. If the remainder equals 1, the account structure conforms to international banking standards. This prevents false alarms on legitimate international wire transfers.",
                "formula_header": "Mathematical Formula:",
                "formula_lines": [
                    "Rearranged_String = IBAN[4:] + IBAN[:4]",
                    "Numeric_Value mod 97 = 1"
                ],
                "algo_header": "Basic Algorithm:",
                "steps": [
                    "1. Ingest candidate IBAN strings from the transaction stream.",
                    "2. Move the first four characters to the end of the string.",
                    "3. Convert letters A through Z to their numerical equivalents (A=10, B=11, ..., Z=35).",
                    "4. Compute the modulo 97 of the resulting numeric integer.",
                    "5. If the remainder equals 1, confirm the string as an authentic international account."
                ]
            },
            {
                "title": "4.3.3 Reverse-Offset Masking Algorithm",
                "desc": "When redacting multiple sensitive numbers in a single paragraph, replacing an entity changes the text length. If replacements are made from left to right, all subsequent character positions shift and become invalid. SecureFlowAI sorts all detected entities in descending order by start index and replaces them from right to left. This ensures that earlier character positions remain completely stable throughout the replacement process.",
                "formula_header": "Mathematical Formula:",
                "formula_lines": [
                    'Token = "[MASKED_" + Type + "_" + HMAC_SHA256(Key, Plaintext)[:8] + "]"'
                ],
                "algo_header": "Basic Algorithm:",
                "steps": [
                    "1. Identify all sensitive entity spans (start index, end index, entity type) within the document.",
                    "2. Sort detected entities in descending order based on their starting character position.",
                    "3. For each entity, generate a unique HMAC token and encrypt the plaintext into the local SQLite vault.",
                    "4. Replace the entity span with the masked token from right to left.",
                    "5. Forward the sanitized text with preserved formatting."
                ]
            },
            {
                "title": "4.3.4 Sliding-Window Transformer Inference",
                "desc": "To inspect long documents that exceed the model's 256-word limit, text is divided into 256-word windows with 50 words of overlap between consecutive windows. The overlap prevents sensitive numbers from being split across window boundaries. If any window contains restricted data, the entire document is flagged as Restricted.",
                "formula_header": "Mathematical Formula:",
                "formula_lines": [
                    "Stride = Window_Size - Overlap = 256 - 50 = 206",
                    "Total_Windows = ceil((N - Overlap) / Stride)"
                ],
                "algo_header": "Basic Algorithm:",
                "steps": [
                    "1. Ingest document text and tokenize the stream into an array of words of length N.",
                    "2. Configure a sliding window of 256 tokens and a stride of 206 tokens.",
                    "3. Slice overlapping windows across the document sequence.",
                    "4. Execute INT8 neural classification on each window.",
                    "5. If any window contains restricted financial data, classify the entire document as Restricted; otherwise, classify as Allowed."
                ]
            }
        ]

        curr_p_xml = ref_p._p
        def insert_paragraph_after(text="", bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=6, font_size=12, is_box=False):
            nonlocal curr_p_xml
            new_p_elem = parse_xml(f'<w:p {nsdecls("w")}/>')
            curr_p_xml.addnext(new_p_elem)
            p_obj = docx.text.paragraph.Paragraph(new_p_elem, doc)
            if text:
                r = p_obj.add_run(text)
                r.font.name = 'Times New Roman'
                r.font.size = Pt(font_size)
                r.bold = bold
                r.italic = italic
                if is_box:
                    r.font.name = 'Consolas'
                    r.font.size = Pt(10)
            p_obj.alignment = align
            p_obj.paragraph_format.line_spacing = 1.15
            p_obj.paragraph_format.space_before = Pt(space_before)
            p_obj.paragraph_format.space_after = Pt(space_after)
            curr_p_xml = new_p_elem
            return p_obj

        for sec in sections:
            # Subsection Title
            insert_paragraph_after(sec["title"], bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=12, space_after=4, font_size=12)
            # Description text
            insert_paragraph_after(sec["desc"], align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=6, font_size=12)
            # Formula Header
            insert_paragraph_after(sec["formula_header"], bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=4, space_after=2, font_size=11.5)
            # Formula lines (centered or indented)
            for fl in sec["formula_lines"]:
                insert_paragraph_after(fl, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=1, space_after=2, font_size=11)
            # Algo Header
            insert_paragraph_after(sec["algo_header"], bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=4, space_after=2, font_size=11.5)
            # Steps
            for st in sec["steps"]:
                insert_paragraph_after(st, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=1, space_after=2, font_size=11.5)

        print("Successfully injected all 4 structured algorithms with formulas and steps into Section 4.3.")

    # =========================================================================
    # 3. SAVE TO ALL TARGET PATHS
    # =========================================================================
    targets = [
        "SecureFlowAI_Final_College_Project_Report.docx",
        "SecureFlowAI_Final_College_Report.docx",
        "docs/reports/SecureFlowAI_Final_College_Project_Report.docx",
        "docs/reports/SecureFlowAI_Final_College_Report.docx"
    ]
    for target in targets:
        try:
            doc.save(target)
            print(f"SUCCESS: Saved updated report to {target} ({os.path.getsize(target)/1024:.1f} KB)")
        except Exception as e:
            print(f"Error saving to {target}: {e}")

if __name__ == "__main__":
    update_report()
